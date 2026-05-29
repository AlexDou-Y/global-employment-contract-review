#!/usr/bin/env python3
"""Extract plain text from contract files for review.

Supports .txt, .md, .docx, and text-based .pdf files. Scanned PDFs may require
OCR outside this script.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path


for stream in (sys.stdout, sys.stderr):
    if hasattr(stream, "reconfigure"):
        stream.reconfigure(encoding="utf-8")


def extract_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def extract_docx(path: Path) -> str:
    try:
        import docx  # type: ignore
    except ImportError as exc:
        raise SystemExit("python-docx is required for .docx extraction") from exc

    document = docx.Document(str(path))
    parts: list[str] = []
    parts.extend(p.text for p in document.paragraphs if p.text.strip())
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_pdf(path: Path) -> str:
    try:
        import pdfplumber  # type: ignore
    except ImportError as exc:
        raise SystemExit("pdfplumber is required for .pdf extraction") from exc

    parts: list[str] = []
    with pdfplumber.open(str(path)) as pdf:
        for index, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                parts.append(f"\n--- Page {index} ---\n{text}")
    return "\n".join(parts).strip()


def extract(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return extract_txt(path)
    if suffix == ".docx":
        return extract_docx(path)
    if suffix == ".pdf":
        return extract_pdf(path)
    raise SystemExit(f"Unsupported file type: {suffix}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract plain text from an employment contract file.")
    parser.add_argument("input", type=Path, help="Input .docx, .pdf, .txt, or .md file")
    parser.add_argument("-o", "--output", type=Path, help="Optional output .txt path")
    args = parser.parse_args()

    text = extract(args.input)
    if args.output:
        args.output.write_text(text, encoding="utf-8")
    else:
        print(text)


if __name__ == "__main__":
    main()
